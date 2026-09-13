# -*- coding: utf-8 -*-
"""Small spreadsheet-upload helpers used only by Dežurstva staff import."""
from io import BytesIO


def _upload_bytes(request, name):
    upload = request.form.get(name)
    if upload is None:
        return b'', ''
    filename = getattr(upload, 'filename', '') or getattr(upload, 'name', '') or ''
    reader = getattr(upload, 'read', None)
    if callable(reader):
        data = reader()
    elif getattr(upload, 'file', None) is not None:
        data = upload.file.read()
    else:
        data = bytes(upload)
    if isinstance(data, str):
        data = data.encode('utf-8')
    return data or b'', str(filename)


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float):
        return str(value)
    return str(value)


def _workbook_rows(data, filename):
    lower = filename.lower()
    if lower.endswith('.xlsx') or lower.endswith('.xlsm'):
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    if lower.endswith('.xls'):
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data)
        sheet = workbook.sheet_by_index(0)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]
    raise ValueError('Podprte so datoteke .xls, .xlsx in .xlsm.')
